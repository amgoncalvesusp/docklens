from __future__ import annotations

from dataclasses import replace

import openpyxl

from docklens import export
from docklens.export_views import (
    detail_dataframe,
    input_qc_dataframe,
    key_residue_coverage_dataframe,
    residue_matrix_dataframe,
    summary_dataframe,
)


def _identity_result(multi_source_result):
    summaries = tuple(
        replace(
            summary,
            source_pose_label=f"dock{summary.pose}",
            source_molecule_name=f"{summary.ligand_id}|dock{summary.pose}",
            score_type="Gold.Goldscore.Fitness",
        )
        for summary in multi_source_result.summaries
    )
    details = tuple(
        replace(
            detail,
            source_pose_label=f"dock{detail.pose}",
            source_molecule_name=f"{detail.ligand_id}|dock{detail.pose}",
            score_type="Gold.Goldscore.Fitness",
        )
        for detail in multi_source_result.details
    )
    qc = tuple(
        replace(
            record,
            ligand_id="LIG-A" if record.source_id == "S000001" else "LIG-B",
            source_pose_label="dock1",
            source_molecule_name="LIG-A|dock1",
            docking_score=-8.0,
            score_type="Gold.Goldscore.Fitness",
        )
        for record in multi_source_result.input_qc
    )
    return replace(
        multi_source_result,
        summaries=summaries,
        details=details,
        input_qc=qc,
    )


def test_identity_is_repeated_in_all_analytical_views(multi_source_result):
    result = _identity_result(multi_source_result)

    summary = summary_dataframe(result)
    detail = detail_dataframe(result)
    coverage = key_residue_coverage_dataframe(result)
    matrix = residue_matrix_dataframe(result)
    qc = input_qc_dataframe(result)

    for frame in (summary, detail, coverage, qc):
        assert {
            "ligand_id",
            "source_pose_label",
            "source_molecule_name",
            "docking_score",
            "score_type",
        } <= set(frame.columns)
    assert {
        ("Identity", "ligand_id"),
        ("Identity", "source_pose_label"),
        ("Identity", "source_molecule_name"),
        ("Identity", "docking_score"),
        ("Identity", "score_type"),
    } <= set(matrix.columns)
    assert summary.loc[0, "source_pose_label"] == "dock1"
    assert detail.loc[0, "source_molecule_name"] == "LIG-A|dock1"


def test_xlsx_contains_ligand_index_and_consistent_pose_ids(
    multi_source_result, tmp_path
):
    result = _identity_result(multi_source_result)
    output = export.export_xlsx(result, tmp_path / "identity.xlsx")
    workbook = openpyxl.load_workbook(output, read_only=True)

    assert "Ligand Index" in workbook.sheetnames
    index = workbook["Ligand Index"]
    headers = [cell.value for cell in index[1]]
    assert headers == [
        "ligand_id",
        "source_pose_label",
        "source_molecule_name",
        "docking_score",
        "score_type",
        "source_file",
        "pose",
        "pose_id",
        "source_id",
        "resolution_method",
        "source_path",
    ]
    assert index.max_row == len(result.summaries) + 1
    assert {
        row[7]
        for row in index.iter_rows(min_row=2, values_only=True)
    } == {summary.pose_id for summary in result.summaries}
    workbook.close()
